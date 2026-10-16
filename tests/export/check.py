"""Checks an exported LungMarker workbook against Daniel's master-sheet rules (status.md section 3).
Usage: python3 tests/export/check.py tests/export/out   (reads out.xlsx + out.state.json)"""
import json, re, sys, zipfile

def zipfile_serial(stem):
    x = zipfile.ZipFile(f"{stem}.xlsx").read("xl/worksheets/sheet1.xml").decode()
    m = re.search(r'<c r="B2"[^>]*><v>([^<]+)</v>', x)
    return m.group(1) if m else "?"
from openpyxl import load_workbook

stem = sys.argv[1]
wb = load_workbook(f"{stem}.xlsx")
state = json.load(open(f"{stem}.state.json"))
ws, tr = wb["Measurements"], wb["Edit trail"]

COLS = ["Patient ID", "Scan date", "Timepoint", "Series #", "Slice thickness (mm)", "Nodule ID", "Lobe", "Slice #", "Type",
        "Long axis (mm)", "Short axis (mm)", "Mean diameter (mm)", "Solid component (mm)", "Volume (mm3)", "Match to prior",
        "Minutes on scan", "Measured by", "Notes"]
header = [c.value for c in ws[1]]
rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
nf_rows = [r for r in rows if r["Match to prior"] == "Not found"]
rows = [r for r in rows if r["Match to prior"] != "Not found"]
cells = {(c.row, c.column_letter): c for r in ws.iter_rows(min_row=2, max_col=18) for c in r}
nod = {(n["engine_scan_id"], n["nodule_label"]): n for n in state["nodules"]}
sid = {("Baseline"): "RIDER-TEST_test", ("Follow-up 1"): "RIDER-TEST_retest"}
key = lambda r: (sid[r["Timepoint"]], r["Nodule ID"])
results = []

def check(name, ok, evidence):
    results.append((name, ok, evidence))

def show(v):
    return "(empty)" if v is None else repr(v)

# 1 columns
check("18 columns, exact names and order", header == COLS,
      f"{len(header)} columns: " + " | ".join(map(str, header)))

# 2 formats
bad = []
for r in rows:
    pass
    if not re.fullmatch(r"Baseline|Follow-up \d+", str(r["Timepoint"])): bad.append(f"timepoint {show(r['Timepoint'])}")
    if r["Lobe"] is not None and r["Lobe"] not in {"RUL", "RML", "RLL", "LUL", "LLL"}: bad.append(f"lobe {show(r['Lobe'])}")
    if not re.fullmatch(r"N\d+", str(r["Nodule ID"])): bad.append(f"nodule id {show(r['Nodule ID'])}")
    if not re.fullmatch(r"[A-Z]{2,3}", str(r["Measured by"])): bad.append(f"measured by {show(r['Measured by'])}")
check("Formats like his example row (Timepoint, Lobe, N1, initials)", not bad,
      "; ".join(bad) or "; ".join(f"{r['Timepoint']} {r['Nodule ID']} {show(r['Lobe'])} {r['Measured by']}" for r in rows))
date_cells = [ws.cell(row=i, column=2) for i in range(2, ws.max_row + 1)]
bad = [f"B{c.row} type={c.data_type} format={c.number_format} value={c.value!r}" for c in date_cells if not (c.is_date and c.number_format == "yyyy-mm-dd")]
check("Scan date is a real Excel date, shown yyyy-mm-dd", not bad, "; ".join(bad) or
      f"B2: stored as date {date_cells[0].value:%Y-%m-%d} (Excel number {zipfile_serial(stem)}), format '{date_cells[0].number_format}' → displays 2006-01-10")

# 3 mean
bad = [f"{r['Timepoint']} {r['Nodule ID']}: ({r['Long axis (mm)']}+{r['Short axis (mm)']})/2 = {(r['Long axis (mm)'] + r['Short axis (mm)']) / 2:.2f} but sheet says {r['Mean diameter (mm)']}"
       for r in rows if abs((r["Long axis (mm)"] + r["Short axis (mm)"]) / 2 - r["Mean diameter (mm)"]) > 0.051]
check("Mean diameter = (long + short) / 2", not bad, "; ".join(bad) or "; ".join(
    f"{r['Nodule ID']} ({r['Timepoint'][:4]}): ({r['Long axis (mm)']}+{r['Short axis (mm)']})/2 → {r['Mean diameter (mm)']}" for r in rows))

# 4 slice #
bad, ok_ev = [], []
for r in rows:
    n = nod[key(r)]
    expected = n["ai_draft"].get("key_slice_number") if n["source"] == "AI" else n["slice_number"]
    (bad if r["Slice #"] != expected else ok_ev).append(f"{r['Nodule ID']} ({r['Timepoint'][:4]}): {r['Slice #']} vs key slice {expected}")
check("Slice # = viewer image number of the key (largest) slice", not bad, "; ".join(bad or ok_ev))

# 5 volume: confirmed outline -> volume filled; not confirmed -> empty
bad, ev = [], []
for r in rows:
    n = nod[key(r)]
    got = r["Volume (mm3)"]
    if bool(n.get("outline_confirmed")) != (got is not None): bad.append(f"{r['Nodule ID']} ({r['Timepoint'][:4]}) confirmed={n.get('outline_confirmed')} volume={show(got)}")
    ai = n["ai_draft"].get("volume")
    ev.append(f"{r['Nodule ID']} ({r['Timepoint'][:4]}): confirmed={bool(n.get('outline_confirmed'))} → {show(got)}" + (f" (AI {ai}, recomputed after hand edit)" if got is not None and ai is not None and got != ai else ""))
check("Volume filled for every confirmed outline, empty otherwise", not bad, "; ".join(bad or ev))

# 6 match to prior
init = state["afterInit"]
ai_filled = [a for a in init if a["match_to_prior"]]
bad = [f"AI pre-filled {a['label']} = {a['match_to_prior']}" for a in ai_filled]
for r in rows:
    if r["Timepoint"] == "Baseline" and r["Match to prior"] is not None: bad.append(f"baseline {r['Nodule ID']} has {show(r['Match to prior'])}")
    if r["Timepoint"] != "Baseline" and not (re.fullmatch(r"N\d+", str(r["Match to prior"])) or r["Match to prior"] == "N/A"): bad.append(f"follow-up {r['Nodule ID']} has {show(r['Match to prior'])}")
check("Match to prior: Daniel's pick or N/A, never AI-filled, empty on Baseline", not bad,
      "; ".join(bad) or "after AI draft: " + ", ".join(f"{a['label']}={show(a['match_to_prior'] or None)} (AI suggests {a['ai_suggests'] or 'nothing'})" for a in init)
      + " → exported: " + ", ".join(f"{r['Nodule ID']} ({r['Timepoint'][:4]})={show(r['Match to prior'])}" for r in rows))

# 7 solid component
bad = [f"{r['Nodule ID']} ({r['Timepoint'][:4]}) {r['Type']} solid={show(r['Solid component (mm)'])}" for r in rows
       if (r["Type"] != "Part-solid" and r["Solid component (mm)"] is not None) or (r["Type"] == "Part-solid" and r["Solid component (mm)"] is None)]
check("Solid component empty for Solid/Ground-glass, filled for Part-solid", not bad,
      "; ".join(bad) or "; ".join(f"{r['Nodule ID']} ({r['Timepoint'][:4]}) {r['Type']} → {show(r['Solid component (mm)'])}" for r in rows))

# 8 minutes
scans = {s["engine_scan_id"]: s for s in state["scans"]}
bad, ev = [], []
for r in rows:
    s = scans[sid[r["Timepoint"]]]
    exp = -(-s["total_seconds"] // 60)
    (bad if r["Minutes on scan"] != exp or s.get("active_minutes") != -(-s["active_seconds"] // 60) else ev).append(
        f"{r['Timepoint']}: open → approved {s['total_seconds']} s → {r['Minutes on scan']} min exported (active {s['active_minutes']} min kept for Results)")
check("Minutes on scan = total time open → approved (Daniel's definition); active kept too", not bad, "; ".join(dict.fromkeys(bad or ev)))

# 9 blanks: read the raw sheet XML (openpyxl turns an empty-text cell into None and would hide it)
import zipfile
xml = zipfile.ZipFile(f"{stem}.xlsx").read("xl/worksheets/sheet1.xml").decode()
raw = re.findall(r'<c r="([A-Z]+)(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', xml)
bad = [f"{col}{row} {header[ord(col) - 65]} stored as {body or '(stub)'}" for col, row, attr, body in raw
       if row != "1" and (body in ("", "<v></v>", "<v>0</v>") or re.search(r"<v>(null|None)</v>", body or ""))]
blank_expected = sum(1 for r in rows for v in r.values() if v is None)
check("Blank cells truly empty (no '', 0, null in the file)", not bad,
      "; ".join(bad[:8]) + (f" … {len(bad)} cells" if len(bad) > 8 else "") if bad else f"{blank_expected} blank cells, none written into the file")

# 10 >= 4 mm
bad = [f"{r['Timepoint'][:4]} {r['Nodule ID']} long axis {r['Long axis (mm)']} mm" for r in rows if r["Long axis (mm)"] < 4]
check("Only nodules ≥ 4 mm (long axis)", not bad, "; ".join(bad) or "smallest long axis: " + str(min(r["Long axis (mm)"] for r in rows)) + " mm")

# notes: Daniel's text only
bad = [f"{r['Timepoint'][:4]} {r['Nodule ID']}: {r['Notes']!r}" for r in rows if r["Notes"] is not None and re.search(r"AI-assisted|accepted unchanged|Added by reviewer", r["Notes"])]
check("Notes = Daniel's own text only", not bad, "; ".join(bad) or "; ".join(f"{r['Timepoint'][:4]} {r['Nodule ID']}: {show(r['Notes'])}" for r in rows))

# 11 trail
th = [c.value for c in tr[1]]
trail = [dict(zip(th, [c.value for c in r])) for r in tr.iter_rows(min_row=2)]
fields = {t["Field"] for t in trail}
ai_shown = any(t.get("AI value") not in (None, "") for t in trail)
draft_rows = [t for t in trail if str(t["Field"]).startswith("AI draft")]
approved = [t for t in trail if str(t["Field"]).startswith("scan_approved")]
unchanged = [r for r in rows if "accepted unchanged" in str(r["Notes"])]
covered = all(any(t["Nodule ID"] == r["Nodule ID"] and t["Timepoint"] == r["Timepoint"] and str(t["Field"]).startswith("AI draft") for t in trail) for r in rows)
check("Edit trail: AI draft for every exported nodule", covered,
      f"{len(draft_rows)} 'AI draft' rows; AI value column filled on changed fields: {ai_shown}; nodules accepted unchanged: {len(unchanged)}")
results_rows = [t for t in trail if t["Field"] == "Review result"]
check("Edit trail: review result per exported row (moved from Notes)", len(results_rows) == len(rows) + len(nf_rows),
      "; ".join(f"{t['Timepoint'][:4]} {t['Nodule ID']}: {t['New value']}" for t in results_rows))
check("Edit trail: 'not found' recorded as AI draft + review result", any(t["Field"] == "Review result" and "confirmed not found" in str(t["New value"]) for t in trail)
      and any(t["Field"] == "AI draft" and "not found in this scan" in str(t["AI value"]) for t in trail),
      "; ".join(f"{t['Nodule ID']}: {t['Field']} → {t['AI value'] or t['New value']}" for t in trail if "prior" in str(t["Nodule ID"]) and t["Field"] in ("AI draft", "Review result")))
check("Edit trail: Daniel's changes (old → new, who, when)", any(t["Changed by"] == "DR" and t["New value"] not in (None, "") for t in trail),
      "fields changed: " + ", ".join(sorted(f for f in fields if f)))
central = [t for t in trail if t["Field"] == "central_areas_checked"]
check("Edit trail: central areas (hilum / mediastinum) checked, per scan", len(central) == len(scans) and all(t["New value"] == "yes" and t["Changed by"] == "DR" for t in central),
      "; ".join(f"{t['Timepoint']}: {t['New value']} by {t['Changed by']} at {t['Changed at']}" for t in central) or "none")
check("Edit trail: who approved each scan", len(approved) == len(scans), "; ".join(f"{t['Timepoint']}: {t['New value']} by {t['Changed by']} at {t['Changed at']}" for t in approved))

# prior nodules not found on the follow-up
nfa = state["notFoundAfterInit"]
check("Not-found items created from the engine list (+ hand-added prior), Pending for Daniel", sorted(x["prior"] for x in nfa) == ["N2", "N4"] and all(x["status"] == "Pending" and not x["match_to_prior"] for x in nfa),
      "; ".join(f"prior {x['prior']}: {x['status']}, {'engine not_found_in_current' if x['from_engine'] else 'added by hand on prior'}" for x in nfa) + " · prior N3 (rejected) not asked")
ok = len(nf_rows) == 1 and nf_rows[0]["Nodule ID"] is None and nf_rows[0]["Lobe"] == "LLL" and nf_rows[0]["Timepoint"] == "Follow-up 1" \
    and all(nf_rows[0][c] is None for c in ["Slice #", "Type", "Long axis (mm)", "Short axis (mm)", "Mean diameter (mm)", "Solid component (mm)", "Volume (mm3)"]) \
    and nf_rows[0]["Notes"] == "Not seen on this scan, likely resolved" and nf_rows[0]["Measured by"] == "DR"
check("Confirmed 'not found' → one row: Match to prior = Not found + Daniel's note, no sizes", ok,
      " | ".join(f"{k}={show(v)}" for k, v in (nf_rows[0].items() if nf_rows else [])) if nf_rows else "no row")
check("Prior nodule marked 'it's here' is not exported as Not found", len(nf_rows) == 1, f"{len(nf_rows)} Not found row(s); prior N4 was marked present")

# daily "only rows not exported yet" exports: every row exactly once, changed rows again
daily = state["daily"]
seen = [r for d in daily[:2] for r in d["rows"]]
ok = (sorted(seen) == sorted(set(seen)) and len(seen) == 5 and daily[2]["rows"] == ["Baseline N1"] and daily[3]["rows"] == [])
check("Daily export: no duplicates, changed row re-exported, nothing twice", ok,
      " · ".join(f"{d['day']}: {', '.join(d['rows']) or '(nothing new)'}" for d in daily))
check("Re-exported row is flagged as UPDATED in the Read me sheet", any("UPDATED" in l for l in daily[2]["readme"]) and not any("UPDATED" in l for l in daily[1]["readme"]),
      f"Oct 17 Read me: {' / '.join(daily[2]['readme'])}")

# 11 / 12 / 13: edit-trail rows (Stav, 16 Oct)
import datetime as _dt
th = [c.value for c in tr[1]]
trows = [dict(zip(th, [c.value for c in r])) for r in tr.iter_rows(min_row=2)]
ch_col = th.index("Changed at") + 1
ch_cells = [tr.cell(row=i, column=ch_col) for i in range(2, tr.max_row + 1)]
ch_vals = [c for c in ch_cells if c.value not in (None, "")]
check("Edit trail times = real Excel date-times (local), format yyyy-mm-dd hh:mm", ch_vals and all(isinstance(c.value, _dt.datetime) and c.number_format == "yyyy-mm-dd hh:mm" for c in ch_vals),
      f"{len(ch_vals)} timed rows; e.g. {ch_vals[-1].value} as {ch_vals[-1].number_format}" if ch_vals else "no timed rows")
rm = [c.value for c in wb["Read me"]["A"] if c.value]
tzl = [l for l in rm if "local time" in str(l)]
check("Read me states the time zone of the trail times", bool(tzl) and "UTC" in tzl[0], tzl[0] if tzl else "no time-zone line")
empty = [t for t in trows if t["Field"] not in ("AI draft", "Review result") and t["Previous value"] in (None, "") and t["New value"] in (None, "")]
check("Edit trail has no non-changes (∅ → ∅)", not empty, f"{len(empty)} empty change row(s)" + (": " + ", ".join(f"{t['Nodule ID']} {t['Field']}" for t in empty[:5]) if empty else ""))
chain_bad, chains = [], {}
for t in trows:
    if t["Field"] in ("AI draft", "Review result") or str(t["Field"]).startswith("outline") or not t["Nodule ID"]:
        continue
    k = (t["Timepoint"], t["Nodule ID"], t["Field"])
    if k in chains and str(chains[k] or "") != str(t["Previous value"] or ""):
        chain_bad.append(f"{k[1]} {k[2]}: previous {t['Previous value']!r} but last new value was {chains[k]!r}")
    chains[k] = t["New value"]
multi = sum(1 for k in chains if sum(1 for t in trows if (t["Timepoint"], t["Nodule ID"], t["Field"]) == k) > 1)
check("Repeated edits: each Previous value = the value just before it", not chain_bad, "; ".join(chain_bad) or f"{multi} field(s) edited more than once, all chained correctly")

w = max(len(n) for n, _, _ in results)
for n, ok, ev in results:
    print(f"{'PASS' if ok else 'FAIL'}  {n.ljust(w)}  │ {ev}")
print(f"\n{sum(ok for _, ok, _ in results)}/{len(results)} pass")
